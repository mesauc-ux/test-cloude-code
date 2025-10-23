"""
Ders Programı Uygulaması - Route Handler Modülü
Route handlers module for the Course Schedule Application
"""
from flask import Blueprint, render_template, request, jsonify, send_file, make_response
from datetime import datetime
import io
import json
import random
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from weasyprint import HTML

import database as db

# Blueprint oluştur
main_bp = Blueprint('main', __name__)

# Global değişken - ders programı verisi
schedule_data = None

@main_bp.route('/')
def index():
    """Ana sayfa"""
    return render_template('index.html')

# ============= ÖĞRETMEN İŞLEMLERİ (TEACHER OPERATIONS) =============

@main_bp.route('/teachers', methods=['GET'])
def get_teachers():
    """Tüm öğretmenleri getir"""
    teachers = db.get_teachers()
    return jsonify([dict(teacher) for teacher in teachers])

@main_bp.route('/teachers', methods=['POST'])
def add_teacher():
    """Yeni öğretmen ekle"""
    data = request.get_json()
    teacher_id = db.add_teacher(
        data['name'],
        data['surname'],
        data['branch'],
        data.get('schedule', '{}')
    )
    return jsonify({'success': True, 'id': teacher_id})

@main_bp.route('/teachers/<int:teacher_id>', methods=['PUT'])
def update_teacher(teacher_id):
    """Öğretmen bilgilerini güncelle"""
    data = request.get_json()
    db.update_teacher(
        teacher_id,
        data['name'],
        data['surname'],
        data['branch'],
        data.get('schedule', '{}'),
        data.get('blocked_slots')
    )
    return jsonify({'success': True})

@main_bp.route('/teachers/<int:teacher_id>', methods=['DELETE'])
def delete_teacher(teacher_id):
    """Öğretmeni sil"""
    db.delete_teacher(teacher_id)
    return jsonify({'success': True})

# ============= ÖĞRENCİ İŞLEMLERİ (STUDENT OPERATIONS) =============

@main_bp.route('/students', methods=['GET'])
def get_students():
    """Tüm öğrencileri getir"""
    students = db.get_students()
    return jsonify([dict(student) for student in students])

@main_bp.route('/students', methods=['POST'])
def add_student():
    """Yeni öğrenci ekle"""
    data = request.get_json()
    student_id = db.add_student(
        data['name'],
        data['surname'],
        data['class'],
        data.get('restrictions', ''),
        data.get('priorities', ''),
        data.get('manual_lessons', '')
    )
    return jsonify({'success': True, 'id': student_id})

@main_bp.route('/students/<int:student_id>', methods=['PUT'])
def update_student(student_id):
    """Öğrenci bilgilerini güncelle"""
    data = request.get_json()
    db.update_student(
        student_id,
        data['name'],
        data['surname'],
        data['class'],
        data.get('restrictions', ''),
        data.get('priorities', ''),
        data.get('manual_lessons', ''),
        data.get('teacher_blocks', '')
    )
    return jsonify({'success': True})

@main_bp.route('/students/<int:student_id>', methods=['DELETE'])
def delete_student(student_id):
    """Öğrenciyi sil"""
    db.delete_student(student_id)
    return jsonify({'success': True})

# ============= DERS PROGRAMI İŞLEMLERİ (SCHEDULE OPERATIONS) =============

@main_bp.route('/generate_schedule', methods=['POST'])
def generate_schedule():
    """
    4 haftalık ders programı oluştur

    Özellikleri:
    - Öğretmen ve öğrenci sayısına göre otomatik slot dağılımı
    - Çakışma kontrolü
    - Öğretmen bloke edilmiş saatleri
    - Öğrenci kısıtlamaları
    - Rastgele dağılım ile adil program
    """
    global schedule_data

    try:
        # Öğretmen ve öğrenci verilerini al
        teachers = db.get_teachers()
        students = db.get_students()

        if not teachers or not students:
            return jsonify({
                'success': False,
                'message': 'Lütfen önce öğretmen ve öğrenci ekleyin!'
            })

        # 4 haftalık program yapısı
        weeks = ['Hafta 1', 'Hafta 2', 'Hafta 3', 'Hafta 4']
        days = ['Pazartesi', 'Salı', 'Çarşamba', 'Perşembe', 'Cuma']
        time_slots = [
            '09:00-10:00', '10:00-11:00', '11:00-12:00', '12:00-13:00',
            '13:00-14:00', '14:00-15:00', '15:00-16:00', '16:00-17:00'
        ]

        # Program verisi
        schedule = {}
        conflicts = []

        # Her öğrenci için ders sayısını hesapla (öğretmen sayısına göre)
        lessons_per_student = len(teachers)
        total_slots = len(weeks) * len(days) * len(time_slots)

        # Her öğrenci için ders ata
        for student in students:
            student_id = str(student['id'])
            student_name = f"{student['name']} {student['surname']}"
            student_class = student['class']

            # Öğrenci kısıtlamalarını parse et
            restrictions = json.loads(student.get('restrictions', '{}')) if student.get('restrictions') else {}
            manual_lessons = json.loads(student.get('manual_lessons', '{}')) if student.get('manual_lessons') else {}

            # Her öğretmen için ders ata
            for teacher in teachers:
                teacher_id = str(teacher['id'])
                teacher_name = f"{teacher['name']} {teacher['surname']}"
                teacher_branch = teacher['branch']

                # Öğretmen bloke edilmiş slotları
                blocked_slots = json.loads(teacher.get('blocked_slots', '{}')) if teacher.get('blocked_slots') else {}

                # Rastgele hafta, gün ve saat seç
                assigned = False
                attempts = 0
                max_attempts = 50

                while not assigned and attempts < max_attempts:
                    week = random.choice(weeks)
                    day = random.choice(days)
                    time_slot = random.choice(time_slots)

                    # Slot anahtarı oluştur
                    slot_key = f"{week}|{day}|{time_slot}"

                    # Çakışma kontrolü
                    conflict = False

                    # 1. Öğretmen bloke kontrolü
                    if slot_key in blocked_slots:
                        conflict = True

                    # 2. Öğrenci kısıtlama kontrolü
                    if slot_key in restrictions:
                        conflict = True

                    # 3. Schedule'da bu slot zaten kullanılmış mı?
                    if slot_key in schedule:
                        # Aynı öğretmen farklı öğrenciye ders veremez (aynı anda)
                        for entry in schedule[slot_key]:
                            if entry['teacher_id'] == teacher_id:
                                conflict = True
                                break
                            # Aynı öğrenci aynı anda farklı öğretmenle ders alamaz
                            if entry['student_id'] == student_id:
                                conflict = True
                                break

                    if not conflict:
                        # Slotu ata
                        if slot_key not in schedule:
                            schedule[slot_key] = []

                        schedule[slot_key].append({
                            'teacher_id': teacher_id,
                            'teacher_name': teacher_name,
                            'teacher_branch': teacher_branch,
                            'student_id': student_id,
                            'student_name': student_name,
                            'student_class': student_class,
                            'lesson': teacher_branch
                        })

                        assigned = True

                    attempts += 1

                if not assigned:
                    conflicts.append({
                        'type': 'assignment_failed',
                        'teacher': teacher_name,
                        'student': student_name,
                        'message': f'{teacher_name} için {student_name} öğrencisine ders atanamadı'
                    })

        # Başarılı program oluşturuldu
        schedule_data = {
            'success': True,
            'schedule': schedule,
            'conflicts': conflicts,
            'stats': {
                'total_lessons': sum(len(v) for v in schedule.values()),
                'total_slots': len(schedule),
                'teachers_count': len(teachers),
                'students_count': len(students),
                'conflicts_count': len(conflicts)
            },
            'weeks': weeks,
            'days': days,
            'time_slots': time_slots
        }

        return jsonify(schedule_data)

    except Exception as e:
        return jsonify({
            'success': False,
            'message': f'Program oluşturulurken hata: {str(e)}'
        })

@main_bp.route('/update_schedule', methods=['POST'])
def update_schedule():
    """Ders programını güncelle (drag-drop için)"""
    global schedule_data

    data = request.get_json()
    # TODO: Schedule güncelleme mantığını ekle

    return jsonify({'success': True})

# ============= KAYDETME İŞLEMLERİ (SAVE OPERATIONS) =============

@main_bp.route('/save_schedule', methods=['POST'])
def save_schedule():
    """Ders programını veritabanına kaydet"""
    global schedule_data

    data = request.get_json()
    schedule_name = data.get('name', f'Program_{datetime.now().strftime("%Y%m%d_%H%M%S")}')

    # Öğretmen ve öğrenci snapshot'larını al
    teachers = db.get_teachers()
    students = db.get_students()

    teachers_snapshot = json.dumps([dict(t) for t in teachers])
    students_snapshot = json.dumps([dict(s) for s in students])

    schedule_id = db.save_schedule(
        schedule_name,
        json.dumps(schedule_data),
        teachers_snapshot,
        students_snapshot
    )

    return jsonify({'success': True, 'id': schedule_id})

@main_bp.route('/saved_schedules', methods=['GET'])
def get_saved_schedules():
    """Kaydedilmiş tüm programları getir"""
    schedules = db.get_saved_schedules()
    return jsonify([dict(s) for s in schedules])

@main_bp.route('/saved_schedules/<int:schedule_id>', methods=['GET'])
def load_saved_schedule(schedule_id):
    """Kaydedilmiş programı yükle"""
    global schedule_data

    schedule = db.get_saved_schedule_by_id(schedule_id)
    if schedule:
        schedule_data = json.loads(schedule['schedule_data'])
        return jsonify({
            'success': True,
            'schedule': schedule_data,
            'teachers': json.loads(schedule['teachers_snapshot']) if schedule['teachers_snapshot'] else [],
            'students': json.loads(schedule['students_snapshot']) if schedule['students_snapshot'] else []
        })
    return jsonify({'success': False, 'message': 'Program bulunamadı'})

@main_bp.route('/saved_schedules/<int:schedule_id>', methods=['DELETE'])
def delete_saved_schedule(schedule_id):
    """Kaydedilmiş programı sil"""
    db.delete_saved_schedule(schedule_id)
    return jsonify({'success': True})

# ============= EXPORT İŞLEMLERİ (EXPORT OPERATIONS) =============

@main_bp.route('/export/excel', methods=['POST'])
def export_excel():
    """
    Ders programını Excel olarak indir
    - Her hafta için ayrı sayfa
    - Renkli ve formatlı hücreler
    - Öğretmen ve öğrenci bilgileri
    """
    global schedule_data

    if not schedule_data or not schedule_data.get('success'):
        # Boş program için uyarı
        wb = Workbook()
        ws = wb.active
        ws.title = "Uyarı"
        ws['A1'] = "Lütfen önce bir ders programı oluşturun!"
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return send_file(
            output,
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            as_attachment=True,
            download_name=f'ders_programi_{datetime.now().strftime("%Y%m%d_%H%M%S")}.xlsx'
        )

    # Excel workbook oluştur
    wb = Workbook()
    wb.remove(wb.active)  # Default sheet'i sil

    # Stil tanımlamaları
    header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=12)
    time_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    time_font = Font(bold=True, size=10)
    cell_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    lesson_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
    center_alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

    weeks = schedule_data.get('weeks', ['Hafta 1', 'Hafta 2', 'Hafta 3', 'Hafta 4'])
    days = schedule_data.get('days', ['Pazartesi', 'Salı', 'Çarşamba', 'Perşembe', 'Cuma'])
    time_slots = schedule_data.get('time_slots', [])
    schedule = schedule_data.get('schedule', {})

    # Her hafta için bir sayfa oluştur
    for week in weeks:
        ws = wb.create_sheet(title=week)

        # Başlık satırı
        ws.merge_cells('A1:F1')
        title_cell = ws['A1']
        title_cell.value = f"DERS PROGRAMI - {week}"
        title_cell.font = Font(bold=True, size=14, color="FFFFFF")
        title_cell.fill = PatternFill(start_color="203864", end_color="203864", fill_type="solid")
        title_cell.alignment = center_alignment
        ws.row_dimensions[1].height = 30

        # Kolon başlıkları (Günler)
        ws['A2'] = 'Saat'
        ws['A2'].fill = header_fill
        ws['A2'].font = header_font
        ws['A2'].alignment = center_alignment
        ws['A2'].border = cell_border
        ws.column_dimensions['A'].width = 15

        for col_idx, day in enumerate(days, start=2):
            cell = ws.cell(row=2, column=col_idx)
            cell.value = day
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = center_alignment
            cell.border = cell_border
            ws.column_dimensions[cell.column_letter].width = 25

        # Zaman slotları ve dersler
        row_idx = 3
        for time_slot in time_slots:
            ws.cell(row=row_idx, column=1, value=time_slot)
            ws.cell(row=row_idx, column=1).fill = time_fill
            ws.cell(row=row_idx, column=1).font = time_font
            ws.cell(row=row_idx, column=1).alignment = center_alignment
            ws.cell(row=row_idx, column=1).border = cell_border
            ws.row_dimensions[row_idx].height = 60

            for col_idx, day in enumerate(days, start=2):
                cell = ws.cell(row=row_idx, column=col_idx)
                slot_key = f"{week}|{day}|{time_slot}"

                if slot_key in schedule:
                    lessons = schedule[slot_key]
                    lesson_texts = []
                    for lesson in lessons:
                        lesson_text = f"{lesson['teacher_name']}\n{lesson['lesson']}\n({lesson['student_name']})"
                        lesson_texts.append(lesson_text)

                    cell.value = "\n---\n".join(lesson_texts)
                    cell.fill = lesson_fill
                else:
                    cell.value = ""

                cell.alignment = center_alignment
                cell.border = cell_border

            row_idx += 1

    # Excel dosyasını kaydet
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    return send_file(
        output,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        as_attachment=True,
        download_name=f'ders_programi_{datetime.now().strftime("%Y%m%d_%H%M%S")}.xlsx'
    )

@main_bp.route('/export/pdf', methods=['POST'])
def export_pdf():
    """
    Ders programını PDF olarak indir
    - Her hafta için tablo
    - Profesyonel görünüm
    - Yazdırma için optimize edilmiş
    """
    global schedule_data

    if not schedule_data or not schedule_data.get('success'):
        html_content = """
        <!DOCTYPE html>
        <html>
        <head>
            <meta charset="UTF-8">
            <style>
                body { font-family: Arial, sans-serif; padding: 20px; }
                .warning { color: red; font-size: 18px; }
            </style>
        </head>
        <body>
            <p class="warning">Lütfen önce bir ders programı oluşturun!</p>
        </body>
        </html>
        """
        pdf = HTML(string=html_content).write_pdf()
        return send_file(
            io.BytesIO(pdf),
            mimetype='application/pdf',
            as_attachment=True,
            download_name=f'ders_programi_{datetime.now().strftime("%Y%m%d_%H%M%S")}.pdf'
        )

    weeks = schedule_data.get('weeks', [])
    days = schedule_data.get('days', [])
    time_slots = schedule_data.get('time_slots', [])
    schedule = schedule_data.get('schedule', {})
    stats = schedule_data.get('stats', {})

    # HTML içeriği oluştur
    html_content = f"""
    <!DOCTYPE html>
    <html>
    <head>
        <meta charset="UTF-8">
        <title>Ders Programı</title>
        <style>
            @page {{
                size: A4 landscape;
                margin: 1cm;
            }}
            body {{
                font-family: 'Arial', sans-serif;
                font-size: 10px;
            }}
            .header {{
                text-align: center;
                background-color: #203864;
                color: white;
                padding: 15px;
                margin-bottom: 20px;
            }}
            .header h1 {{
                margin: 0;
                font-size: 24px;
            }}
            .stats {{
                background-color: #f0f0f0;
                padding: 10px;
                margin-bottom: 20px;
                border-radius: 5px;
            }}
            .stats-grid {{
                display: grid;
                grid-template-columns: repeat(3, 1fr);
                gap: 10px;
            }}
            .stat-item {{
                text-align: center;
            }}
            .stat-value {{
                font-size: 20px;
                font-weight: bold;
                color: #366092;
            }}
            .week-section {{
                page-break-after: always;
                margin-bottom: 30px;
            }}
            .week-section:last-child {{
                page-break-after: avoid;
            }}
            .week-title {{
                background-color: #366092;
                color: white;
                padding: 10px;
                font-size: 18px;
                font-weight: bold;
                text-align: center;
                margin-bottom: 10px;
            }}
            table {{
                width: 100%;
                border-collapse: collapse;
                margin-bottom: 20px;
            }}
            th {{
                background-color: #366092;
                color: white;
                padding: 8px;
                text-align: center;
                border: 1px solid #ddd;
                font-size: 11px;
            }}
            td {{
                border: 1px solid #ddd;
                padding: 8px;
                text-align: center;
                vertical-align: top;
                min-height: 60px;
            }}
            .time-col {{
                background-color: #D9E1F2;
                font-weight: bold;
                width: 80px;
            }}
            .lesson {{
                background-color: #E2EFDA;
                padding: 5px;
                margin: 2px 0;
                border-radius: 3px;
                font-size: 9px;
            }}
            .lesson-teacher {{
                font-weight: bold;
                color: #203864;
            }}
            .lesson-subject {{
                color: #366092;
            }}
            .lesson-student {{
                color: #666;
                font-style: italic;
            }}
            .footer {{
                text-align: center;
                margin-top: 20px;
                font-size: 9px;
                color: #666;
            }}
        </style>
    </head>
    <body>
        <div class="header">
            <h1>DERS PROGRAMI YÖNETİM SİSTEMİ</h1>
            <p>4 Haftalık Ders Programı</p>
            <p>Oluşturulma Tarihi: {datetime.now().strftime("%d.%m.%Y %H:%M")}</p>
        </div>

        <div class="stats">
            <div class="stats-grid">
                <div class="stat-item">
                    <div class="stat-value">{stats.get('teachers_count', 0)}</div>
                    <div>Öğretmen</div>
                </div>
                <div class="stat-item">
                    <div class="stat-value">{stats.get('students_count', 0)}</div>
                    <div>Öğrenci</div>
                </div>
                <div class="stat-item">
                    <div class="stat-value">{stats.get('total_lessons', 0)}</div>
                    <div>Toplam Ders</div>
                </div>
            </div>
        </div>
    """

    # Her hafta için tablo oluştur
    for week in weeks:
        html_content += f"""
        <div class="week-section">
            <div class="week-title">{week}</div>
            <table>
                <thead>
                    <tr>
                        <th class="time-col">Saat</th>
        """

        for day in days:
            html_content += f"<th>{day}</th>"

        html_content += """
                    </tr>
                </thead>
                <tbody>
        """

        for time_slot in time_slots:
            html_content += f"""
                    <tr>
                        <td class="time-col">{time_slot}</td>
            """

            for day in days:
                slot_key = f"{week}|{day}|{time_slot}"
                html_content += "<td>"

                if slot_key in schedule:
                    lessons = schedule[slot_key]
                    for lesson in lessons:
                        html_content += f"""
                        <div class="lesson">
                            <div class="lesson-teacher">{lesson['teacher_name']}</div>
                            <div class="lesson-subject">{lesson['lesson']}</div>
                            <div class="lesson-student">({lesson['student_name']})</div>
                        </div>
                        """

                html_content += "</td>"

            html_content += "</tr>"

        html_content += """
                </tbody>
            </table>
        </div>
        """

    html_content += """
        <div class="footer">
            <p>© 2024 Ders Programı Yönetim Sistemi | Modüler Flask Uygulaması</p>
        </div>
    </body>
    </html>
    """

    # PDF oluştur
    pdf = HTML(string=html_content).write_pdf()

    return send_file(
        io.BytesIO(pdf),
        mimetype='application/pdf',
        as_attachment=True,
        download_name=f'ders_programi_{datetime.now().strftime("%Y%m%d_%H%M%S")}.pdf'
    )

@main_bp.route('/export/html', methods=['POST'])
def export_html():
    """
    Ders programını HTML olarak indir
    - Tarayıcıda görüntülenebilir
    - Responsive tasarım
    - Yazdırma desteği
    """
    global schedule_data

    if not schedule_data or not schedule_data.get('success'):
        html_content = """
        <!DOCTYPE html>
        <html>
        <head>
            <meta charset="UTF-8">
            <title>Ders Programı</title>
            <style>
                body { font-family: Arial, sans-serif; padding: 20px; text-align: center; }
                .warning { color: red; font-size: 18px; }
            </style>
        </head>
        <body>
            <p class="warning">Lütfen önce bir ders programı oluşturun!</p>
        </body>
        </html>
        """
        return send_file(
            io.BytesIO(html_content.encode('utf-8')),
            mimetype='text/html',
            as_attachment=True,
            download_name=f'ders_programi_{datetime.now().strftime("%Y%m%d_%H%M%S")}.html'
        )

    weeks = schedule_data.get('weeks', [])
    days = schedule_data.get('days', [])
    time_slots = schedule_data.get('time_slots', [])
    schedule = schedule_data.get('schedule', {})
    stats = schedule_data.get('stats', {})

    # HTML içeriği oluştur
    html_content = f"""
    <!DOCTYPE html>
    <html lang="tr">
    <head>
        <meta charset="UTF-8">
        <meta name="viewport" content="width=device-width, initial-scale=1.0">
        <title>Ders Programı - {datetime.now().strftime("%d.%m.%Y")}</title>
        <style>
            * {{
                margin: 0;
                padding: 0;
                box-sizing: border-box;
            }}

            body {{
                font-family: 'Segoe UI', Tahoma, Geneva, Verdana, sans-serif;
                background: linear-gradient(135deg, #667eea 0%, #764ba2 100%);
                padding: 20px;
            }}

            .container {{
                max-width: 1400px;
                margin: 0 auto;
                background: white;
                border-radius: 15px;
                box-shadow: 0 20px 60px rgba(0,0,0,0.3);
                overflow: hidden;
            }}

            .header {{
                background: linear-gradient(135deg, #203864 0%, #366092 100%);
                color: white;
                padding: 30px;
                text-align: center;
            }}

            .header h1 {{
                font-size: 32px;
                margin-bottom: 10px;
            }}

            .header p {{
                font-size: 16px;
                opacity: 0.9;
            }}

            .stats {{
                display: grid;
                grid-template-columns: repeat(auto-fit, minmax(200px, 1fr));
                gap: 20px;
                padding: 30px;
                background: #f8f9fa;
            }}

            .stat-card {{
                background: white;
                padding: 20px;
                border-radius: 10px;
                text-align: center;
                box-shadow: 0 2px 10px rgba(0,0,0,0.1);
            }}

            .stat-value {{
                font-size: 36px;
                font-weight: bold;
                color: #366092;
                margin-bottom: 5px;
            }}

            .stat-label {{
                color: #666;
                font-size: 14px;
            }}

            .content {{
                padding: 30px;
            }}

            .week-section {{
                margin-bottom: 40px;
                page-break-inside: avoid;
            }}

            .week-title {{
                background: #366092;
                color: white;
                padding: 15px;
                font-size: 24px;
                font-weight: bold;
                text-align: center;
                border-radius: 10px 10px 0 0;
            }}

            table {{
                width: 100%;
                border-collapse: collapse;
                box-shadow: 0 2px 10px rgba(0,0,0,0.1);
                border-radius: 0 0 10px 10px;
                overflow: hidden;
            }}

            th {{
                background-color: #4a5568;
                color: white;
                padding: 15px 10px;
                text-align: center;
                font-size: 14px;
                font-weight: 600;
            }}

            td {{
                border: 1px solid #e2e8f0;
                padding: 12px 8px;
                text-align: center;
                vertical-align: top;
                min-height: 80px;
                background: white;
            }}

            .time-col {{
                background-color: #edf2f7;
                font-weight: bold;
                color: #2d3748;
                width: 100px;
            }}

            .lesson {{
                background: linear-gradient(135deg, #E2EFDA 0%, #C5E1A5 100%);
                padding: 8px;
                margin: 4px 0;
                border-radius: 6px;
                border-left: 4px solid #7CB342;
                font-size: 12px;
                transition: transform 0.2s;
            }}

            .lesson:hover {{
                transform: translateY(-2px);
                box-shadow: 0 4px 8px rgba(0,0,0,0.1);
            }}

            .lesson-teacher {{
                font-weight: bold;
                color: #203864;
                margin-bottom: 3px;
            }}

            .lesson-subject {{
                color: #366092;
                font-weight: 600;
                margin-bottom: 3px;
            }}

            .lesson-student {{
                color: #666;
                font-style: italic;
                font-size: 11px;
            }}

            .footer {{
                text-align: center;
                padding: 20px;
                background: #f8f9fa;
                color: #666;
                font-size: 14px;
            }}

            @media print {{
                body {{
                    background: white;
                    padding: 0;
                }}
                .container {{
                    box-shadow: none;
                }}
                .week-section {{
                    page-break-after: always;
                }}
                .week-section:last-child {{
                    page-break-after: avoid;
                }}
                .lesson:hover {{
                    transform: none;
                }}
            }}

            @media (max-width: 768px) {{
                .header h1 {{
                    font-size: 24px;
                }}
                .stats {{
                    grid-template-columns: 1fr;
                }}
                table {{
                    font-size: 11px;
                }}
                th, td {{
                    padding: 8px 4px;
                }}
            }}
        </style>
    </head>
    <body>
        <div class="container">
            <div class="header">
                <h1>📚 DERS PROGRAMI YÖNETİM SİSTEMİ</h1>
                <p>4 Haftalık Ders Programı</p>
                <p>Oluşturulma Tarihi: {datetime.now().strftime("%d.%m.%Y %H:%M")}</p>
            </div>

            <div class="stats">
                <div class="stat-card">
                    <div class="stat-value">{stats.get('teachers_count', 0)}</div>
                    <div class="stat-label">👨‍🏫 Öğretmen</div>
                </div>
                <div class="stat-card">
                    <div class="stat-value">{stats.get('students_count', 0)}</div>
                    <div class="stat-label">👨‍🎓 Öğrenci</div>
                </div>
                <div class="stat-card">
                    <div class="stat-value">{stats.get('total_lessons', 0)}</div>
                    <div class="stat-label">📖 Toplam Ders</div>
                </div>
                <div class="stat-card">
                    <div class="stat-value">{stats.get('total_slots', 0)}</div>
                    <div class="stat-label">🕐 Kullanılan Slot</div>
                </div>
            </div>

            <div class="content">
    """

    # Her hafta için tablo oluştur
    for week in weeks:
        html_content += f"""
                <div class="week-section">
                    <div class="week-title">{week}</div>
                    <table>
                        <thead>
                            <tr>
                                <th class="time-col">⏰ Saat</th>
        """

        for day in days:
            html_content += f"<th>{day}</th>"

        html_content += """
                            </tr>
                        </thead>
                        <tbody>
        """

        for time_slot in time_slots:
            html_content += f"""
                            <tr>
                                <td class="time-col">{time_slot}</td>
            """

            for day in days:
                slot_key = f"{week}|{day}|{time_slot}"
                html_content += "<td>"

                if slot_key in schedule:
                    lessons = schedule[slot_key]
                    for lesson in lessons:
                        html_content += f"""
                        <div class="lesson">
                            <div class="lesson-teacher">{lesson['teacher_name']}</div>
                            <div class="lesson-subject">{lesson['lesson']}</div>
                            <div class="lesson-student">({lesson['student_name']})</div>
                        </div>
                        """

                html_content += "</td>"

            html_content += "</tr>"

        html_content += """
                        </tbody>
                    </table>
                </div>
        """

    html_content += """
            </div>

            <div class="footer">
                <p>© 2024 Ders Programı Yönetim Sistemi | Modüler Flask Uygulaması</p>
                <p>Bu program otomatik olarak oluşturulmuştur.</p>
            </div>
        </div>
    </body>
    </html>
    """

    return send_file(
        io.BytesIO(html_content.encode('utf-8')),
        mimetype='text/html',
        as_attachment=True,
        download_name=f'ders_programi_{datetime.now().strftime("%Y%m%d_%H%M%S")}.html'
    )
